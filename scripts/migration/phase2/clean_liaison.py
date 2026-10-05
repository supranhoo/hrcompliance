import openpyxl, re, datetime, collections, difflib
from clean_common import *

ESIC_WB = 'b6051f46-1788441813413_Liaison_with_ESIC.xlsx'; ESIC_N = 'Liaison_with_ESIC.xlsx'
GOVT_WB = 'cb910a6c-1788442341975_Liaisoning_with_Govt_Officials.xlsx'; GOVT_N = 'Liaisoning_with_Govt_Officials.xlsx'
ESIC_COLS = [('source_workbook','ok'),('source_sheet','ok'),('source_row','ok'),('flag_codes','ok'),('sl_no','ok'),('information_received_on_clean','ok'),('work_completed_on_clean','ok'),
             ('tat_days_source','ok'),('tat_days_recomputed_calendar','ok'),('tat_matches','ok'),('particulars','text'),('particulars_char_count','ok'),('contains_long_digit_token','ok'),('possible_overlap_with','ok')]
GOVT_COLS = [('source_workbook','ok'),('source_sheet','ok'),('source_row','ok'),('flag_codes','ok'),('sr_no','ok'),('sr_cell_kind','ok'),('month_start_clean','ok'),('case_detail','text'),('case_detail_char_count','ok'),
             ('remark_kind','ok'),('remark_as_date_clean','ok'),('remark_text','text'),('possible_overlap_with','ok')]
LONGDIGIT = re.compile(r'\b\d{9,}\b')

def norm_key(s): return re.sub(r'[^a-z0-9]', '', (s or '').lower())

def run(pseudo):
    E = Area('LIAISON_ESIC', ESIC_COLS); G = Area('LIAISON_GOVT', GOVT_COLS)
    ev = openpyxl.load_workbook(SRC + ESIC_WB, data_only=True).active; ef = openpyxl.load_workbook(SRC + ESIC_WB).active
    gv = openpyxl.load_workbook(SRC + GOVT_WB, data_only=True).active; gf = openpyxl.load_workbook(SRC + GOVT_WB).active
    ESH, GSH = ev.title, gv.title
    erows = []; grows = []
    for r in range(3, ev.max_row + 1):
        v = [ev.cell(r, c).value for c in range(1, 6)]
        if any(x is not None for x in v): erows.append((r, v))
    for r in range(3, gv.max_row + 1):
        v = [gv.cell(r, c).value for c in range(1, 5)]
        if any(x is not None for x in v): grows.append((r, v))
    # overlap detection (no merging)
    pairs = []
    for er, ev_ in erows:
        for gr, gv_ in grows:
            a, b = norm_key(ev_[2])[:140], norm_key(gv_[2])[:140]
            ratio = difflib.SequenceMatcher(None, a, b).ratio()
            if ratio >= 0.85: pairs.append((er, gr, round(ratio, 3)))
    eo = {er: (gr, ra) for er, gr, ra in pairs}; go = {gr: (er, ra) for er, gr, ra in pairs}
    # ESIC
    for r, v in erows:
        flags = []
        p = ws_norm(v[2]) if isinstance(v[2], str) else v[2]
        if isinstance(v[2], str) and p != v[2]: flags.append('WS_NORMALIZED'); E.log_tf(ESIC_N, ESH, r, 'particulars', 'WS_NORMALIZE', '[free text]', '[free text]', sens='text', note='trailing/repeated whitespace only')
        b, d = v[1], v[3]
        bi, di = iso(b), iso(d)
        if not isinstance(b, datetime.datetime): flags.append('RECEIVED_DATE_NOT_A_DATE')
        if not isinstance(d, datetime.datetime): flags.append('COMPLETED_DATE_NOT_A_DATE')
        rec = (d - b).days if isinstance(b, datetime.datetime) and isinstance(d, datetime.datetime) else None
        match = (rec == v[4]) if rec is not None else None
        if match is False: flags.append('TAT_MISMATCH')
        if rec is not None and rec < 0: flags.append('COMPLETED_BEFORE_RECEIVED')
        long = bool(LONGDIGIT.search(p or ''))
        if long: flags.append('CONTAINS_LONG_DIGIT_TOKEN_SENSITIVE')
        if r in eo: flags.append('POSSIBLE_OVERLAP_WITH_GOVT_LOG')
        E.cleaned.append(dict(source_workbook=ESIC_N, source_sheet=ESH, source_row=r, flag_codes=';'.join(flags) or None, sl_no=v[0], information_received_on_clean=bi, work_completed_on_clean=di, tat_days_source=v[4],
                              tat_days_recomputed_calendar=rec, tat_matches=match, particulars=p, particulars_char_count=len(p or ''), contains_long_digit_token=long, possible_overlap_with=f'{GOVT_N}:{GSH}:row {eo[r][0]} (similarity {eo[r][1]})' if r in eo else None))
        E.log_tf(ESIC_N, ESH, r, 'dates', 'DATE_TO_ISO', f'{b}|{d}', f'{bi}|{di}')
    # GOVT
    for r, v in grows:
        flags = []
        cd = ws_norm(v[2]) if isinstance(v[2], str) else v[2]
        if isinstance(v[2], str) and cd != v[2]: flags.append('WS_NORMALIZED'); G.log_tf(GOVT_N, GSH, r, 'case_detail', 'WS_NORMALIZE', '[free text]', '[free text]', sens='text', note='trailing/repeated whitespace only')
        kind = 'formula' if isinstance(gf.cell(r, 1).value, str) and str(gf.cell(r, 1).value).startswith('=') else 'hardcoded'
        if kind == 'hardcoded': flags.append('SR_HARDCODED_NOT_FORMULA')
        mo = v[1]
        if not isinstance(mo, datetime.datetime): flags.append('MONTH_NOT_A_DATE')
        rem = v[3]
        if isinstance(rem, datetime.datetime): rk, rd_, rt = 'DATE', iso(rem), None
        elif isinstance(rem, str):
            rt = ws_norm(rem); rk, rd_ = 'TEXT', None
            if rt != rem: flags.append('REMARK_WS_NORMALIZED')
        else: rk, rd_, rt = 'BLANK', None, None; flags.append('REMARK_BLANK')
        if cd and len(cd) < 8: flags.append('VERY_SHORT_CASE_DETAIL_PLACEHOLDER_LIKE'); G.review.append(dict(category='PLACEHOLDER_LIKE_ENTRY', source_workbook=GOVT_N, source_sheet=GSH, source_row=r, original_value='[short case detail]', detail='case detail shorter than 8 characters', proposed_action='owner to confirm whether a real case'))
        if r in go: flags.append('POSSIBLE_OVERLAP_WITH_ESIC_LOG')
        G.cleaned.append(dict(source_workbook=GOVT_N, source_sheet=GSH, source_row=r, flag_codes=';'.join(flags) or None, sr_no=v[0] if not isinstance(v[0], str) else v[0], sr_cell_kind=kind, month_start_clean=iso(mo), case_detail=cd,
                              case_detail_char_count=len(cd or ''), remark_kind=rk, remark_as_date_clean=rd_, remark_text=rt, possible_overlap_with=f'{ESIC_N}:{ESH}:row {go[r][0]} (similarity {go[r][1]})' if r in go else None))
        G.log_tf(GOVT_N, GSH, r, 'month/remark', 'DATE_TO_ISO_AND_SPLIT_REMARK', f'{mo}|{type(rem).__name__}', f'{iso(mo)}|{rk}', note='Remark split into date or text by cell type only; the date is NOT interpreted as a completion date')
    # overlaps
    for er, gr, ra in pairs:
        for A_, a_wb, a_sh, a_row, o_wb, o_sh, o_row in ((E, ESIC_N, ESH, er, GOVT_N, GSH, gr), (G, GOVT_N, GSH, gr, ESIC_N, ESH, er)):
            A_.duplicates.append(dict(duplicate_type='POSSIBLE_OVERLAPPING_EVENT_ACROSS_LOGS', source_workbook=a_wb, source_sheet=a_sh, source_row=a_row, group_key=f'similarity {ra}', group_rows=f'{o_wb}:{o_sh}:row {o_row}', resolution='NOT MERGED - both kept; owner to choose single source of truth'))
        for A_ in (E, G): A_.review.append(dict(category='OVERLAPPING_EVENT', source_workbook=f'{ESIC_N} row {er} / {GOVT_N} row {gr}', source_sheet='-', source_row=f'{er}/{gr}', original_value='[free text withheld]', detail=f'text similarity {ra}', proposed_action='owner to decide single source of truth'))
    E.review.append(dict(category='SENSITIVE_MEDICAL_PERSONAL_NARRATIVE', source_workbook=ESIC_N, source_sheet=ESH, source_row=f'all {len(erows)} rows', original_value='[withheld]', detail='free-text particulars hold personal and medical information; not prepared for DEV', proposed_action='owner decision on access, redaction, retention (OC-21)'))
    G.review.append(dict(category='MONTH_ONLY_DATES', source_workbook=GOVT_N, source_sheet=GSH, source_row=f'all {len(grows)} rows', original_value='first-of-month dates', detail='day of the event is not recorded; month_start_clean is the first of the month', proposed_action='owner: is a day-level date available?'))
    G.review.append(dict(category='REMARK_MIXED_DATE_AND_TEXT', source_workbook=GOVT_N, source_sheet=GSH, source_row=f"{sum(1 for x in G.cleaned if x['remark_kind']=='DATE')} date / {sum(1 for x in G.cleaned if x['remark_kind']=='TEXT')} text / {sum(1 for x in G.cleaned if x['remark_kind']=='BLANK')} blank", original_value='-', detail='Remark column mixes dates and outcomes; meaning of the date not assumed', proposed_action='owner (OC-23)'))
    for A_, rows, name in ((E, erows, 'esic'), (G, grows, 'govt')):
        cnt = collections.Counter()
        for rd in A_.cleaned:
            for f in (rd['flag_codes'] or '').split(';'):
                if f: cnt[f] += 1
        A_.summary += [('cleaned_rows', len(A_.cleaned)), ('rows_flagged_at_least_once', sum(1 for r in A_.cleaned if r['flag_codes'])), ('rejected_rows', len(A_.rejected)), ('overlap_pairs_detected', len(pairs)), ('owner_review_items', len(A_.review)), ('transformations_logged', len(A_.log))]
        A_.summary += [(f'flag_{k}', n) for k, n in sorted(cnt.items())]
    E.summary.append(('esic_rows_with_long_digit_token_flagged', sum(1 for r in E.cleaned if r['contains_long_digit_token'])))
    E.rejected.append(dict(note='no rows rejected: no empty, decorative or formula-only rows in the data region'))
    G.rejected.append(dict(note='no rows rejected: no empty, decorative or formula-only rows in the data region'))
    return E, G
