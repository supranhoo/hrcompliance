import openpyxl, re, datetime, collections, os
from clean_common import *

WB = '988197ed-1788633922075_GRC_Register_Final_31082026_JB.xlsx'
WBN = 'GRC_Register_Final_31082026_JB.xlsx'
SH = '29052026'
COLS = [('source_workbook','ok'),('source_sheet','ok'),('source_row','ok'),('record_class','ok'),('flag_codes','ok'),('sl_no','ok'),('registration_number','ok'),
        ('date_clean','ok'),('date_raw_if_text','ok'),('date_annotation','text'),('employee_name','pii:PERSON'),('employee_id_raw','pii:EMPID'),('employee_id_class','ok'),
        ('designation','freq3'),('dept_raw','ok'),('nature_of_grievance','text'),('responsibility','text'),('action_taken','text'),('closed_on_clean','ok'),
        ('status_raw','ok'),('status_clean','ok'),('remarks','text'),('tat_days_source','ok'),('tat_basis','ok'),('stray_col_O_note','text')]
REG = re.compile(r'^BFCL/GRC/\d{4}/\d{4}-\d{2}$')
IDTOK = re.compile(r'^[A-Za-z]{0,4}\d{3,8}$')

def classify_id(v):
    if v is None or (isinstance(v, str) and not v.strip()): return 'MISSING'
    s = ws_norm_single(str(v))
    toks = [t for t in re.split(r'\s*(?:,|/|&|;|\band\b)\s*', s) if t]
    if len(toks) > 1 and all(IDTOK.match(t) for t in toks): return 'MULTIPLE'
    if len(toks) == 1 and re.fullmatch(r'\d{4,8}', toks[0]): return 'NUMERIC_SINGLE'
    if len(toks) == 1 and IDTOK.match(toks[0]): return 'ALPHANUMERIC_CODE'
    return 'NON_ID_TEXT'

def parse_text_date(s):
    """Return (iso_date_or_None, annotation, status) ; only dd/mm/yyyy where the day > 12 is treated as unambiguous."""
    m = re.findall(r'\b(\d{1,2})/(\d{1,2})/(\d{4})\b', s)
    if len(m) != 1: return None, None, 'AMBIGUOUS_MULTIPLE' if len(m) > 1 else 'UNPARSEABLE'
    d, mo, y = map(int, m[0])
    if d > 12 and 1 <= mo <= 12:
        try: dt = datetime.date(y, mo, d)
        except ValueError: return None, None, 'UNPARSEABLE'
        ann = ws_norm_single(re.sub(r'\b\d{1,2}/\d{1,2}/\d{4}\b', '', s)).strip(' ()')
        return dt.isoformat(), ann or None, 'UNAMBIGUOUS'
    return None, None, 'AMBIGUOUS_DAY_MONTH'

def run(pseudo):
    A = Area('GRC', COLS)
    wv = openpyxl.load_workbook(SRC + WB, data_only=True); wf = openpyxl.load_workbook(SRC + WB)
    ws, wsf = wv[SH], wf[SH]
    asof = ws['K1'].value
    recs, placeholders = [], []
    for r in range(3, ws.max_row + 1):
        v = [ws.cell(r, c).value for c in range(1, 16)]
        if v[1]: recs.append((r, v))
        elif any(x is not None for x in v[:13]) or v[14] is not None:
            if not (isinstance(v[0], str) and not v[0].strip() and all(x is None for x in v[1:])): placeholders.append((r, v))
            else: placeholders.append((r, v))
        elif v[13] is not None: placeholders.append((r, v))
    # --- registration analysis
    regs = collections.defaultdict(list); sls = collections.defaultdict(list)
    for r, v in recs:
        regs[ws_norm_single(v[1])].append(r); sls[v[0]].append(r)
    last_contig = 0
    prev = None
    for r, _ in recs:
        if prev is None or r == prev + 1: last_contig = r
        else: break
        prev = r
    A.summary.append(('source_records_with_registration_number', len(recs)))
    for r, v in recs:
        flags = []
        reg = ws_norm_single(v[1])
        if reg != v[1]: A.log_tf(WBN, SH, r, 'registration_number', 'WS_TRIM', v[1], reg)
        fy = re.search(r'/(\d{4})-(\d{2})$', reg)
        if not REG.match(reg) or (fy and (int(fy.group(1)) + 1) % 100 != int(fy.group(2))): flags.append('MALFORMED_REG')
        if len(regs[reg]) > 1: flags.append('DUP_REG')
        if v[0] is not None and len(sls[v[0]]) > 1: flags.append('SL_DUP')
        m = re.match(r'^BFCL/GRC/(\d{4})/', reg)
        if m and v[0] is not None and int(m.group(1)) != v[0] and 'DUP_REG' not in flags: flags.append('SL_REG_MISMATCH')
        elif m and v[0] is not None and int(m.group(1)) != v[0]: flags.append('SL_REG_MISMATCH')
        if r > last_contig: flags.append('STRAY_ROW_AFTER_PLACEHOLDER_BLOCK')
        # date
        d = v[2]; dclean = None; draw = None; dann = None
        if isinstance(d, datetime.datetime): dclean = iso(d)
        elif isinstance(d, str):
            draw = d; dclean, dann, st = parse_text_date(d)
            if dclean: flags.append('TEXT_DATE_NORMALIZED'); A.log_tf(WBN, SH, r, 'date', 'TEXT_DATE_UNAMBIGUOUS', d, dclean, note='day>12 so dd/mm/yyyy is unambiguous; annotation kept separately')
            else: flags.append('AMBIGUOUS_DATE'); A.review.append(dict(category='AMBIGUOUS_DATE', source_workbook=WBN, source_sheet=SH, source_row=r, original_value=d, detail=st, proposed_action='owner to choose registration date'))
        else: flags.append('DATE_MISSING')
        # employee id
        idc = classify_id(v[4])
        if idc != 'NUMERIC_SINGLE': flags.append('EMP_ID_' + idc)
        if idc in ('MULTIPLE', 'NON_ID_TEXT', 'ALPHANUMERIC_CODE'):
            A.review.append(dict(category='EMP_ID_' + idc, source_workbook=WBN, source_sheet=SH, source_row=r, original_value=v[4], detail='raw value preserved; not split or interpreted', proposed_action='owner ruling'))
        # status
        sraw = v[11]; sclean = None
        if isinstance(sraw, str):
            t = ws_norm_single(sraw)
            sclean = 'Closed' if t.lower() == 'closed' else ('Open' if t.lower() == 'open' else t)
            if sclean != sraw: flags.append('STATUS_NORMALIZED'); A.log_tf(WBN, SH, r, 'status', 'STATUS_CLOSED_CANON' if sclean == 'Closed' else 'WS_TRIM', repr(sraw), sclean)
        else: flags.append('STATUS_BLANK')
        closed = iso(v[10]) if isinstance(v[10], datetime.datetime) else None
        if sclean == 'Open' and closed is None: pass
        elif sclean == 'Closed' and closed is None: flags.append('CLOSED_WITHOUT_DATE')
        tat = v[13]; basis = None
        if tat is not None: basis = 'calendar days, closed date minus date' if closed else ('open item measured to TODAY() cached ' + iso(asof) + ' (volatile)')
        else: flags.append('TAT_MISSING')
        def tx(field, val, single):
            n = ws_norm_single(val) if single else ws_norm(val)
            if isinstance(val, str) and n != val: A.log_tf(WBN, SH, r, field, 'WS_NORMALIZE', val, n, sens='text' if field in ('nature_of_grievance','action_taken','remarks','responsibility') else ('pii:PERSON' if field == 'employee_name' else 'ok'))
            return n
        rowd = dict(source_workbook=WBN, source_sheet=SH, source_row=r, sl_no=v[0], registration_number=reg, date_clean=dclean, date_raw_if_text=draw, date_annotation=dann,
                    employee_name=tx('employee_name', v[3], True), employee_id_raw=(str(v[4]) if v[4] is not None else None), employee_id_class=idc,
                    designation=tx('designation', v[5], True), dept_raw=tx('dept_raw', v[6], True), nature_of_grievance=tx('nature_of_grievance', v[7], False),
                    responsibility=tx('responsibility', v[8], True), action_taken=tx('action_taken', v[9], False), closed_on_clean=closed, status_raw=sraw, status_clean=sclean,
                    remarks=tx('remarks', v[12], False), tat_days_source=tat, tat_basis=basis, stray_col_O_note=tx('stray_col_O_note', v[14], False))
        rowd['_flags'] = flags
        A.cleaned.append(rowd)
    # classification, duplicates and review rows
    for rd in A.cleaned:
        fl = rd.pop('_flags'); rd['flag_codes'] = ';'.join(fl) if fl else None
        crit = [f for f in fl if f in ('MALFORMED_REG','DUP_REG','SL_DUP','SL_REG_MISMATCH','AMBIGUOUS_DATE','STRAY_ROW_AFTER_PLACEHOLDER_BLOCK','STATUS_BLANK','CLOSED_WITHOUT_DATE','DATE_MISSING','EMP_ID_MULTIPLE','EMP_ID_NON_ID_TEXT','EMP_ID_ALPHANUMERIC_CODE')]
        rd['record_class'] = 'FLAGGED' if (crit or fl) else 'CLEAN'
    for reg, rows in regs.items():
        if len(rows) > 1:
            for r in rows: A.duplicates.append(dict(duplicate_type='DUPLICATE_REGISTRATION_NUMBER', source_workbook=WBN, source_sheet=SH, source_row=r, group_key=reg, group_rows=','.join(map(str, rows)), resolution='NOT RESOLVED - both kept; owner ruling required'))
            A.review.append(dict(category='DUPLICATE_REGISTRATION_NUMBER', source_workbook=WBN, source_sheet=SH, source_row=','.join(map(str, rows)), original_value=reg, detail='same registration number on different rows', proposed_action='owner to assign/confirm numbers (no auto-fix)'))
    for r, v in recs:
        fy2 = re.search(r'/(\d{4})-(\d{2})$', ws_norm_single(v[1]))
        if not REG.match(ws_norm_single(v[1])) or (fy2 and (int(fy2.group(1)) + 1) % 100 != int(fy2.group(2))):
            A.review.append(dict(category='MALFORMED_REGISTRATION_NUMBER', source_workbook=WBN, source_sheet=SH, source_row=r, original_value=v[1], detail='does not match BFCL/GRC/NNNN/YYYY-YY with consecutive financial years', proposed_action='owner to confirm correct number'))
    for r in [r for r, v in recs if r > last_contig]:
        A.review.append(dict(category='STRAY_ROW', source_workbook=WBN, source_sheet=SH, source_row=r, original_value='(record after placeholder block)', detail='kept in Cleaned_Data flagged; registration also duplicates an earlier record', proposed_action='owner ruling'))
    for rd in A.cleaned:
        if rd['employee_id_class'] == 'MISSING':
            A.review.append(dict(category='EMP_ID_MISSING', source_workbook=WBN, source_sheet=SH, source_row=rd['source_row'], original_value=None, detail='no employee id', proposed_action='owner: acceptable for group/contract records?'))
    # placeholders (formula-only rows) / decorative
    for r, v in placeholders:
        reason = 'FORMULA_PLACEHOLDER_ROW' if (all(x is None for x in v[:13]) and v[14] is None and v[13] is not None) else 'DECORATIVE_OR_WHITESPACE_ROW'
        A.rejected.append(dict(reject_reason=reason, source_workbook=WBN, source_sheet=SH, source_row=r, cached_tat_value=v[13], formula=str(wsf.cell(r, 14).value)[:40] if wsf.cell(r, 14).value else None, raw_cell_A=repr(v[0]) if v[0] is not None else None))
        A.log_tf(WBN, SH, r, 'row', 'REMOVE_FORMULA_ONLY_OR_DECORATIVE_ROW', None, None, note=reason)
    # Sheet1 duplicates
    s1 = wv['Sheet1']; main = {rd['sl_no']: rd for rd in A.cleaned}
    main_raw = {v[0]: (r, v) for r, v in recs}
    for r in range(2, s1.max_row + 1):
        sl = s1.cell(r, 1).value
        if sl is None: continue
        v1 = [s1.cell(r, c).value for c in range(1, 14)]
        mr = main_raw.get(sl)
        same = None
        if mr:
            # Sheet1 columns A..M = Sl, Date, Name, EmpId, Desig, Dept, Nature, Resp, Action, Closed, Status, Remarks, TAT  ->  main cols A, C..N
            pairs = [(v1[0], mr[1][0])] + [(v1[i], mr[1][i + 1]) for i in range(1, 13)]
            same = all(a == b or (isinstance(a, str) and isinstance(b, str) and ws_norm_single(a) == ws_norm_single(b)) for a, b in pairs)
        A.duplicates.append(dict(duplicate_type='SHEET1_DUPLICATE_OF_MAIN_RECORD', source_workbook=WBN, source_sheet='Sheet1', source_row=r, group_key=f'Sl {sl}', group_rows=f'{SH}:{mr[0]}' if mr else None,
                                 resolution='excluded from Cleaned_Data: all 13 comparable fields equal' if same else 'NOT EQUAL - needs owner review'))
        if not same: A.review.append(dict(category='SHEET1_NOT_EXACT_DUPLICATE', source_workbook=WBN, source_sheet='Sheet1', source_row=r, original_value=f'Sl {sl}', detail='differs from main sheet', proposed_action='owner review'))
    # summary
    cnt = collections.Counter(); 
    for rd in A.cleaned:
        for f in (rd['flag_codes'] or '').split(';'):
            if f: cnt[f] += 1
    A.summary += [('cleaned_records_total', len(A.cleaned)), ('records_class_CLEAN', sum(1 for r in A.cleaned if r['record_class'] == 'CLEAN')), ('records_class_FLAGGED', sum(1 for r in A.cleaned if r['record_class'] == 'FLAGGED')),
                  ('rejected_formula_placeholder_or_decorative_rows', len(A.rejected)), ('duplicate_rows_listed', len(A.duplicates)), ('owner_review_items', len(A.review)),
                  ('transformations_logged', len(A.log)), ('registration_numbers_with_duplicates', sum(1 for x in regs.values() if len(x) > 1))]
    A.summary += [(f'flag_{k}', n) for k, n in sorted(cnt.items())]
    A.summary += [('status_values_raw', str(dict(collections.Counter(repr(rd['status_raw']) for rd in A.cleaned)))), ('status_values_clean', str(dict(collections.Counter(rd['status_clean'] for rd in A.cleaned)))),
                  ('employee_id_class_counts', str(dict(collections.Counter(rd['employee_id_class'] for rd in A.cleaned))))]
    return A
