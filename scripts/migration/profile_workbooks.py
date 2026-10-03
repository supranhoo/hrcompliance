#!/usr/bin/env python3
"""Profile every workbook/sheet/column under a directory (default data/source) into JSON + Markdown.

Reports per sheet: dimensions, header row guess, per-column inferred types, null/blank counts, distinct counts,
top values (status-like columns), date range + unparseable/mixed date values, formulas (and cached error values such as
#NAME?/#REF!/#DIV/0!/#N/A/#VALUE!), cross-sheet references, merged cells, exact duplicate rows and duplicate
candidate-key columns. It never modifies source files and never guesses meaning from file names.

Usage: python3 scripts/migration/profile_workbooks.py [source_dir] [out_dir]
"""
import json, re, sys, datetime as dt
from collections import Counter
from pathlib import Path
import openpyxl

ERRORS = {"#NAME?", "#REF!", "#DIV/0!", "#N/A", "#VALUE!", "#NUM!", "#NULL!"}
XREF = re.compile(r"(?:'[^']+'|[A-Za-z0-9_]+)!|\[[^\]]+\]")

def kind(v):
    if v is None or (isinstance(v, str) and not v.strip()): return "blank"
    if isinstance(v, bool): return "bool"
    if isinstance(v, (int, float)): return "number"
    if isinstance(v, (dt.datetime, dt.date)): return "date"
    if isinstance(v, str):
        if v.strip() in ERRORS: return "error"
        return "text"
    return type(v).__name__

def profile_sheet(ws_f, ws_v):
    rows_f = list(ws_f.iter_rows(values_only=True))
    rows_v = list(ws_v.iter_rows(values_only=True)) if ws_v else rows_f
    n = len(rows_f)
    # header guess: first row with >= 50% non-blank text cells among first 10 rows
    hdr_i = 0
    for i, r in enumerate(rows_f[:10]):
        cells = [c for c in r if c is not None]
        if r and len(cells) >= max(2, 0.5 * len(r)) and sum(isinstance(c, str) for c in cells) >= 0.7 * len(cells):
            hdr_i = i; break
    header = [str(c).strip() if c is not None else f"col_{j+1}" for j, c in enumerate(rows_f[hdr_i])] if n else []
    body_f, body_v = rows_f[hdr_i + 1:], rows_v[hdr_i + 1:]
    cols, formulas, errors, xrefs = [], [], [], []
    for j, name in enumerate(header):
        vals = [r[j] if j < len(r) else None for r in body_v]
        kinds = Counter(kind(v) for v in vals)
        nonblank = [v for v in vals if kind(v) != "blank"]
        distinct = Counter(str(v).strip() for v in nonblank)
        c = {"index": j + 1, "name": name, "kinds": dict(kinds), "blank": kinds.get("blank", 0), "distinct": len(distinct)}
        if distinct and len(distinct) <= 25: c["values"] = distinct.most_common(25)
        dates = [v for v in nonblank if isinstance(v, (dt.datetime, dt.date))]
        if dates: c["date_min"], c["date_max"] = str(min(dates)), str(max(dates))
        if dates and kinds.get("text"): c["mixed_date_text"] = True
        dup = [k for k, v in distinct.items() if v > 1]
        if nonblank and len(dup) and len(distinct) > 0.5 * len(nonblank): c["duplicate_values"] = len(dup)
        cols.append(c)
    for r_i, r in enumerate(body_f, start=hdr_i + 2):
        for j, v in enumerate(r):
            if isinstance(v, str) and v.startswith("="):
                f = {"cell": f"{openpyxl.utils.get_column_letter(j+1)}{r_i}", "formula": v[:200]}
                if XREF.search(v): xrefs.append(f)
                formulas.append(f)
    for r_i, r in enumerate(body_v, start=hdr_i + 2):
        for j, v in enumerate(r):
            if isinstance(v, str) and v.strip() in ERRORS:
                errors.append({"cell": f"{openpyxl.utils.get_column_letter(j+1)}{r_i}", "value": v.strip()})
    dup_rows = sum(c - 1 for c in Counter(tuple(r) for r in body_v if any(x is not None for x in r)).values() if c > 1)
    return {"rows_total": n, "header_row": hdr_i + 1, "data_rows": len(body_v), "columns": cols,
            "merged_ranges": [str(m) for m in ws_f.merged_cells.ranges][:50],
            "formula_count": len(formulas), "formula_samples": formulas[:15], "cross_sheet_formulas": xrefs[:15],
            "formula_columns": sorted({re.sub(r"\d+", "", f["cell"]) for f in formulas}),
            "error_cells": errors[:50], "error_count": len(errors), "exact_duplicate_rows": dup_rows}

def main():
    src = Path(sys.argv[1] if len(sys.argv) > 1 else "data/source")
    out = Path(sys.argv[2] if len(sys.argv) > 2 else "docs/migration/profile")
    files = sorted(p for p in src.rglob("*") if p.suffix.lower() in {".xlsx", ".xlsm"})
    if not files:
        print(f"No .xlsx/.xlsm files found under {src}. Nothing profiled."); sys.exit(2)
    out.mkdir(parents=True, exist_ok=True)
    report = {}
    for f in files:
        wb_f = openpyxl.load_workbook(f, data_only=False)
        wb_v = openpyxl.load_workbook(f, data_only=True)
        report[str(f)] = {"sheets": {ws.title: profile_sheet(ws, wb_v[ws.title]) | {"state": ws.sheet_state} for ws in wb_f.worksheets},
                          "defined_names": list(wb_f.defined_names.keys()) if hasattr(wb_f.defined_names, "keys") else []}
    (out / "profile.json").write_text(json.dumps(report, indent=2, default=str))
    md = ["# Workbook profile (generated)\n"]
    for fn, wb in report.items():
        md.append(f"\n## {fn}\n")
        for sn, s in wb["sheets"].items():
            md.append(f"\n### {sn} ({s['state']}) — {s['data_rows']} data rows, header row {s['header_row']}, {s['formula_count']} formulas, {s['error_count']} error cells, {s['exact_duplicate_rows']} duplicate rows\n")
            md.append("| # | Column | Kinds | Blank | Distinct | Notes |\n|---|---|---|---|---|---|")
            for c in s["columns"]:
                note = "; ".join(filter(None, [f"dates {c['date_min']}..{c['date_max']}" if "date_min" in c else "", "MIXED date/text" if c.get("mixed_date_text") else "",
                                               f"values {c['values'][:8]}" if "values" in c else ""]))
                md.append(f"| {c['index']} | {c['name']} | {c['kinds']} | {c['blank']} | {c['distinct']} | {note} |")
    (out / "profile.md").write_text("\n".join(md))
    print(f"Profiled {len(files)} workbook(s) -> {out}")

if __name__ == "__main__":
    main()
