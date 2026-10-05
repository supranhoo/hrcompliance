import openpyxl, collections, datetime, re
f='/tmp/claude-0/src_orig/e8b07389-Contractor_MP_cost_01092026_1_1.xlsx'
wv=openpyxl.load_workbook(f,data_only=True)['Mar 2026']; wf=openpyxl.load_workbook(f)['Mar 2026']
H=[wv.cell(1,c).value for c in range(1,22)]
rows=[]
for r in range(2,wv.max_row+1):
    v=[wv.cell(r,c).value for c in range(1,22)]
    if any(x is not None for x in v): rows.append((r,v))
print('non-empty',len(rows))
mon=collections.Counter()
nomonth=[(r,[str(x)[:25] for x in v if x is not None]) for r,v in rows if not isinstance(v[0],datetime.datetime)]
print('rows without month:',nomonth)
data=[(r,v) for r,v in rows if isinstance(v[0],datetime.datetime)]
print('rows with month',len(data))
# rows with no numeric money at all
def money(v): return [x for x in v[8:18] if isinstance(x,(int,float))]
nomoney=[(r,v) for r,v in data if not isinstance(v[12],(int,float)) and not isinstance(v[16],(int,float))]
print('rows without bill amount/total cost:',len(nomoney)); print(collections.Counter(str(v[7])[:60] for r,v in nomoney).most_common(8))
print('Contractor distinct',len(set(v[1] for r,v in data)))
c=collections.Counter((v[1] or '').strip() for r,v in data)
for k,n in sorted(c.items()): print(f'   {k!r}: {n}',end=';')
print()
for i,n in [(2,'Dept'),(3,'Nature'),(4,'CostCentre'),(5,'LS/PR'),(6,'Unit')]:
    c=collections.Counter(v[i] for r,v in data); print(n,len(c),sorted(c.items(),key=lambda x:-x[1])[:60] if n in('LS/PR','Unit') else len(c))
print('\n---- part 2')
def base(n): return re.sub(r'\s*\(\d+\)\s*$','',(n or '').strip())
bases=collections.Counter(base(v[1]) for r,v in data); print('base contractors',len(bases),sorted(bases.items()))
print('with parenthesised number',sum(1 for r,v in data if re.search(r'\(\d+\)\s*$',v[1] or '')))
print('lead/trail ws in contractor',sum(1 for r,v in data if v[1]!=(v[1] or '').strip()),' dept',sum(1 for r,v in data if isinstance(v[2],str) and v[2]!=v[2].strip()))
# text numbers in numeric columns
numcols={8:'Mandays',9:'Wages',10:'PF',11:'ESI',12:'Bill',13:'PerfInc',14:'Bonus',15:'ProdInc',16:'Total',17:'Hold25',18:'OtherHold'}
for i,n in numcols.items():
    t=collections.Counter(type(v[i]).__name__ for r,v in data); print(f'{n:9}',dict(t), [ (r,v[i]) for r,v in data if isinstance(v[i],str)][:4])
# Bill No patterns
bn=[(r,v[7]) for r,v in data if v[7] is not None and not (isinstance(v[7],str) and 'work' in v[7].lower() or isinstance(v[7],str) and 'bill' in v[7].lower())]
print('bill no types',collections.Counter(type(x).__name__ for r,x in bn)); print('bill no blank rows by month: ',collections.Counter(v[0].strftime('%Y-%m') for r,v in data if v[7] is None and isinstance(v[12],(int,float))).most_common(20))
strbn=[x for r,x in bn if isinstance(x,str)]; print('string bill samples',strbn[:8])
vals=[(str(x).strip()) for r,x in bn]; dup=[k for k,c in collections.Counter(vals).items() if c>1]; print('duplicate Bill No values',len(dup),dup[:12])
# hold amount check: R vs Q*25%
bad=0;tot=0;ex=[]
for r,v in data:
    if isinstance(v[16],(int,float)) and isinstance(v[17],(int,float)):
        tot+=1
        if abs(v[17]-v[16]*0.25)>1: bad+=1; ex.append((r,round(v[16]),round(v[17])))
print('Hold!=25% of total',bad,'of',tot,ex[:6])
# hold release month vs month
rel=collections.Counter()
for r,v in data:
    if isinstance(v[19],datetime.datetime) and isinstance(v[0],datetime.datetime):
        rel[(v[19].year*12+v[19].month)-(v[0].year*12+v[0].month)]+=1
    elif v[19] is not None: rel['nondate:'+str(v[19])[:15]]+=1
print('release - month offset (months)',rel)
# totals inconsistent: bill vs total
print('Total > Bill',sum(1 for r,v in data if isinstance(v[12],(int,float)) and isinstance(v[16],(int,float)) and v[16]>v[12]+1))
# formulas vs hardcoded counts in money columns
fc=collections.Counter(); 
for r,v in data:
    for i in range(8,18):
        x=wf.cell(r,i+1).value
        if isinstance(x,str) and x.startswith('='): fc['formula']+=1
        elif x is not None: fc['hard']+=1
print('money cells formula vs hardcoded',fc)
cons=collections.Counter(str(wf.cell(r,17).value)[:1]+('F' if str(wf.cell(r,17).value).startswith('=') else 'H') for r,v in data); print('Total col formula/hard',cons)
# literal formulas that are constants like =638409
lit=sum(1 for r,v in data for i in range(1,22) if isinstance(wf.cell(r,i).value,str) and re.fullmatch(r'=[\d.+\-*/() ]+',wf.cell(r,i).value))
print('literal-arithmetic formulas (hand-keyed numbers inside formulas)',lit)
# duplicate rows: month+contractor+dept+costcentre
key=collections.Counter((v[0],base(v[1]),v[2],v[4],v[7]) for r,v in data); print('dup month+contractor+dept+CC+bill',[(k[0].strftime('%Y-%m'),k[1],k[2]) for k,c in key.items() if c>1][:8],sum(1 for c in key.values() if c>1))
exact=collections.Counter(tuple(v) for r,v in data); print('exact duplicate rows',sum(c-1 for c in exact.values() if c>1))
# month sum by month
tm=collections.defaultdict(float)
for r,v in data:
    if isinstance(v[16],(int,float)): tm[v[0].strftime('%Y-%m')]+=v[16]
print({k:round(x/1e7,2) for k,x in tm.items()})
print('Mandays text',[ (r,v[8]) for r,v in data if isinstance(v[8],str)][:5])
print('Unit blank',sum(1 for r,v in data if not v[6]),' LS/PR blank',sum(1 for r,v in data if not v[5]),' dept blank',sum(1 for r,v in data if not v[2]),' costcentre blank',sum(1 for r,v in data if not v[4]))
