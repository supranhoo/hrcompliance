import openpyxl, glob, os, zipfile
for f in sorted(glob.glob('/tmp/claude-0/src_orig/*.xlsx')):
    z=zipfile.ZipFile(f); names=z.namelist()
    wb=openpyxl.load_workbook(f)  # formulas
    print('\n###',os.path.basename(f),os.path.getsize(f),'bytes; sheets:',len(wb.sheetnames), '; macros' if any('vba' in n for n in names) else '')
    for ws in wb.worksheets:
        merged=len(ws.merged_cells.ranges)
        hidr=sum(1 for r,d in ws.row_dimensions.items() if d.hidden); hidc=sum(1 for c,d in ws.column_dimensions.items() if d.hidden)
        nform=sum(1 for row in ws.iter_rows() for c in row if isinstance(c.value,str) and c.value.startswith('='))
        print(f"  [{ws.sheet_state}] {ws.title!r} dims={ws.dimensions} max_row={ws.max_row} max_col={ws.max_column} merged={merged} hiddenrows={hidr} hiddencols={hidc} formulas={nform} tables={list(ws.tables)} autofilter={ws.auto_filter.ref} freeze={ws.freeze_panes}")
    print('  defined names:',list(wb.defined_names.keys())[:10], ' external links:', [n for n in names if 'externalLink' in n][:3])
