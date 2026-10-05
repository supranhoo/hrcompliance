import openpyxl, sys, datetime
f,sheet=sys.argv[1],sys.argv[2]; r0=int(sys.argv[3]) if len(sys.argv)>3 else 1; r1=int(sys.argv[4]) if len(sys.argv)>4 else 10**6
wbv=openpyxl.load_workbook(f,data_only=True); wbf=openpyxl.load_workbook(f)
v=wbv[sheet]; fm=wbf[sheet]
for r in range(r0,min(r1,v.max_row)+1):
    out=[]
    for c in range(1,v.max_column+1):
        x=v.cell(r,c).value; fx=fm.cell(r,c).value
        if x is None and fx is None: continue
        s=repr(x) if not isinstance(x,datetime.datetime) else x.strftime('%Y-%m-%d')+('T'+x.strftime('%H:%M') if x.time()!=datetime.time(0) else '')
        if isinstance(fx,str) and fx.startswith('='): s+=' {'+fx[:40]+'}'
        out.append(f"{openpyxl.utils.get_column_letter(c)}={s[:60]}")
    hid=' (HIDDEN)' if fm.row_dimensions[r].hidden else ''
    if out: print(r,hid,' | '.join(out))
