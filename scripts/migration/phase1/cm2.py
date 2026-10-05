import openpyxl, collections, datetime, re
f='/tmp/claude-0/src_orig/e8b07389-Contractor_MP_cost_01092026_1_1.xlsx'
wv=openpyxl.load_workbook(f,data_only=True)
m=wv['Mar 2026']
cons=[]
for r in range(2,m.max_row+1):
    v=[m.cell(r,c).value for c in range(1,22)]
    if isinstance(v[0],datetime.datetime): cons.append((r,v))
# P==Q duplicates by month
dupPQ=collections.Counter(v[0].strftime('%Y-%m') for r,v in cons if isinstance(v[15],(int,float)) and isinstance(v[16],(int,float)) and abs(v[15]-v[16])<0.01)
print('Production-Incentive equals Total (probable column shift) by month',dict(dupPQ))
# exact duplicate rows
ex=collections.defaultdict(list)
for r,v in cons: ex[tuple(v)].append(r)
print('exact duplicate row groups',[(rs, str(k[0].date()), k[1], k[2]) for k,rs in ex.items() if len(rs)>1])
# hidden monthly sheets reconciliation
for name,hdrTotalCol,billCol in [('May 2025',15,11),('June 2025',15,11),('July 2025',15,11),('Aug 2025',16,12),('Sept 2025',16,12)]:
    ws=wv[name]
    rows=[[ws.cell(r,c).value for c in range(1,ws.max_column+1)] for r in range(2,ws.max_row+1)]
    rows=[v for v in rows if isinstance(v[0],datetime.datetime)]
    mo=rows[0][0].strftime('%Y-%m')
    tot=sum(v[hdrTotalCol-1] for v in rows if isinstance(v[hdrTotalCol-1],(int,float)))
    c=[v for r,v in cons if v[0].strftime('%Y-%m')==mo]
    ct=sum(v[16] for v in c if isinstance(v[16],(int,float)))
    print(f'{name}: monthly-sheet rows {len(rows)} total {tot:,.0f}; consolidated rows {len(c)} total {ct:,.0f}; diff {tot-ct:,.0f}')
    # other things in the sheet: non-month rows
    other=[(r+2,[str(x)[:25] for x in v if x is not None][:6]) for r,v in enumerate([[ws.cell(r,c).value for c in range(1,ws.max_column+1)] for r in range(2,ws.max_row+1)]) if any(x is not None for x in v) and not isinstance(v[0],datetime.datetime)]
    if other: print('   non-month rows',other[:4])
    # columns beyond T (CL)
    wide=[(c) for c in range(21,ws.max_column+1) if any(ws.cell(r,c).value is not None for r in range(1,ws.max_row+1))]
    if wide: print('   data beyond col T:',wide[:5])
