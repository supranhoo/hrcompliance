import openpyxl, datetime, collections, re
f='/tmp/claude-0/src_orig/988197ed-1788633922075_GRC_Register_Final_31082026_JB.xlsx'
wv=openpyxl.load_workbook(f,data_only=True); ws=wv['29052026']; wf=openpyxl.load_workbook(f)['29052026']
R=[]
for r in range(3,ws.max_row+1):
    v=[ws.cell(r,c).value for c in range(1,16)]
    if v[1]: R.append((r,v))
print('records',len(R))
sl=[v[0] for r,v in R]; print('Sl dup',[k for k,c in collections.Counter(sl).items() if c>1]); 
reg=[v[1] for r,v in R]; print('Reg dup',[(k,[r for r,v in R if v[1]==k]) for k,c in collections.Counter(reg).items() if c>1])
print('Sl gaps',sorted(set(range(1,max(sl)+1))-set(sl)))
print('Sl not in file order',[ (r,v[0]) for (r,v),(r2,v2) in zip(R,R[1:]) if v2[0]!=v[0]+1][:10])
# registration pattern
pat=collections.Counter(re.sub(r'\d','9',x) for x in reg); print('reg patterns',pat)
# reg number vs Sl
mis=[(r,v[0],v[1]) for r,v in R if not v[1].split('/')[2].lstrip('0')==str(v[0])]; print('reg-vs-Sl mismatch',mis)
# FY of registration vs date
def fy(d): return f"{d.year}-{str((d.year+1))[2:]}" if d.month>=4 else f"{d.year-1}-{str(d.year)[2:]}"
bad=[]
for r,v in R:
    if isinstance(v[2],datetime.datetime):
        if v[1].split('/')[3]!=fy(v[2]): bad.append((r,v[1],v[2].date()))
print('reg FY vs date FY mismatch',bad)
print('Date as str:',[(r,v[2]) for r,v in R if isinstance(v[2],str)])
print('Date range',min(v[2] for r,v in R if isinstance(v[2],datetime.datetime)).date(),max(v[2] for r,v in R if isinstance(v[2],datetime.datetime)).date())
print('Dates out of order',[(r,v[2].date()) for (r,v),(r2,v2) in zip(R,R[1:]) if isinstance(v[2],datetime.datetime) and isinstance(v2[2],datetime.datetime) and v2[2]<v[2]][:10])
print('Status values',collections.Counter(v[11] for r,v in R))
print('Closed-on blank vs status',[(r,v[11],v[10]) for r,v in R if v[10] is None])
print('Status Closed but no date', [r for r,v in R if v[11] and v[11].strip().lower()=='closed' and v[10] is None])
print('closed<date',[(r,v[2].date(),v[10].date()) for r,v in R if isinstance(v[2],datetime.datetime) and isinstance(v[10],datetime.datetime) and v[10]<v[2]])
print('Emp id str',[(r,v[4]) for r,v in R if isinstance(v[4],str)])
print('Emp id dup names w/ different ids / same id different names:')
byid=collections.defaultdict(set)
for r,v in R:
    if v[4] is not None: byid[str(v[4]).strip()].add((v[3] or '').strip())
print('  ids with >1 name',{k:list(x)[:3] for k,x in byid.items() if len(x)>1})
bynm=collections.defaultdict(set)
for r,v in R:
    if v[4] is not None and v[3]: bynm[re.sub(r'[^a-z]','',v[3].lower().replace('mr','').replace('mrs',''))].add(str(v[4]).strip())
print('  names with >1 id',{k:x for k,x in bynm.items() if len(x)>1})
print('Id lengths',collections.Counter(len(str(v[4])) for r,v in R if v[4] is not None))
print('Dept distinct',len(set(v[6] for r,v in R if v[6])));
d=collections.Counter((v[6] or '').strip() for r,v in R); print(sorted(d.items(),key=lambda x:-x[1]))
des=collections.Counter((v[5] or '').strip() for r,v in R); print('Designation',sorted(des.items(),key=lambda x:-x[1])[:12])
resp=collections.Counter((v[8] or '').strip() for r,v in R); print('Resp',sorted(resp.items(),key=lambda x:-x[1])[:25])
for i,n in [(3,'name'),(5,'desig'),(6,'dept'),(8,'resp'),(11,'status')]:
    print('ws lead/trail',n,sum(1 for r,v in R if isinstance(v[i],str) and v[i]!=v[i].strip()))
print('N (TAT) cached vs recomputed mismatch:')
asof=wv['29052026']['K1'].value; print(' K1 cached TODAY()',asof)
mm=[]
for r,v in R:
    d,c,t=v[2],v[10],v[13]
    if isinstance(d,datetime.datetime):
        exp=(c-d).days if isinstance(c,datetime.datetime) else (asof-d).days
        if t!=exp: mm.append((r,t,exp))
print(' ',mm)
print('TAT formula variants',collections.Counter(re.sub(r'\d+','#',str(wf.cell(r,14).value)) for r,v in R))
print('Sl formula variants',collections.Counter('formula' if str(wf.cell(r,1).value).startswith('=') else 'hardcoded' for r,v in R))
print('Remarks col O', [(r,v[14]) for r,v in R if v[14]])
print('Rows 3..144 region: rows >144 records', [r for r,v in R if r>144])
print('Cell errors:',sum(1 for row in ws.iter_rows() for c in row if isinstance(c.value,str) and c.value.startswith('#')))
